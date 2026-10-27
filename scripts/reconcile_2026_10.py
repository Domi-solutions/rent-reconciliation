"""One-off reconciliation, October 2026: bring Mowin's money records up to date
before the new leases, then let a person approve every match.

What it does (in one transaction):
  1. Adds the Family Bank statements given (text PDFs, read by the app's parser;
     the balance check must pass). Rows already in the system are skipped.
     Never supersedes other statements — the upload page's overlap rule would
     retire statements from other banks.
  2. Adds National Bank rows that were checked by eye against the scanned copy
     (CSV: only the Mowin lines; the scanner's own reading was unreliable).
  3. Adds the missing monthly rent + service charges (default May–Oct 2026) for
     every occupied unit with a tenant on record, the same way "Generate
     charges" does (due the 5th of that month, per the lease), then applies any
     unallocated credit.
  4. review: writes a spreadsheet of every unmatched payment with a suggested
     unit. apply: reads the reviewed spreadsheet back and assigns the payments
     whose DECISION column names a unit ("IGNORE" marks a line as not Mowin).

No SMS is sent: these are historical payments.

  # On a COPY — builds everything and writes the sheet for review:
  DATABASE_PATH=copy.db ./venv/bin/python scripts/reconcile_2026_10.py review \\
      --family may_jun.pdf sep_oct.pdf --nbk verified_nbk.csv --out review.xlsx
  # Dry run against the live database (rolls back), then for real:
  DATABASE_PATH=/data/rent.db python scripts/reconcile_2026_10.py apply \\
      --family ... --nbk ... --review reviewed.xlsx [--execute]
"""

import argparse
import csv
import os
import sys
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook                      # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill            # noqa: E402

from src.database.db import (allocate_payment, apply_unallocated_credit,  # noqa: E402
                             generate_id, get_connection)
from src.parsers.pdf_parser import parse_bank_statement              # noqa: E402
from src.reconciliation.unassigned import get_unassigned_credits     # noqa: E402

ACTOR = 'admin:LINCKS (reconciliation Oct 2026)'
DEFAULT_PERIODS = ['2026-05', '2026-06', '2026-07', '2026-08', '2026-09', '2026-10']
MONTHS = {m: i + 1 for i, m in enumerate('JAN FEB MAR APR MAY JUN JUL AUG SEP OCT NOV DEC'.split())}


class DryRun(Exception):
    pass


def _iso(day, raw=''):
    """'02-MAY-2026' → '2026-05-02'; ISO passes through. A row whose year wrapped
    onto a later line ('25-MAY') takes the first 20xx year in its own text."""
    day = (day or '').strip().upper().rstrip('-')
    if len(day) == 10 and day[4] == '-':
        return day
    parts = day.split('-')
    if len(parts) == 2:
        import re
        year = re.search(r'\b(20\d{2})\b', raw or '')
        if not year:
            raise SystemExit(f'Cannot tell the year of {day!r}: {raw[:80]!r}')
        parts.append(year.group(1))
    d, mon, y = parts
    return f'{int(y):04d}-{MONTHS[mon]:02d}-{int(d):02d}'


def _norm(ref):
    return (ref or '').strip().upper()


def _known_refs(conn, org_id):
    return {_norm(r[0]) for r in conn.execute("""
        SELECT bt.mpesa_ref FROM bank_transactions bt JOIN bank_statements bs ON bs.id = bt.statement_id
        WHERE bs.org_id = ? AND bt.mpesa_ref IS NOT NULL""", (org_id,))}


def add_family_statement(conn, org_id, pdf_path):
    result = parse_bank_statement(pdf_path)
    v = result.get('validation') or {}
    if result.get('statement_format') != 'family_bank' or not v.get('valid'):
        raise SystemExit(f'{pdf_path}: not a Family Bank statement that balances ({v}) — stopping.')
    txns = [t for t in result['transactions'] if t.txn_type != 'PARSE_ERROR']
    dates = sorted(_iso(t.transaction_date, t.raw_text) for t in txns)
    known = _known_refs(conn, org_id)
    stmt_id = generate_id('STMT')
    conn.execute("""
        INSERT INTO bank_statements (id, org_id, property_id, filename, file_path, period_start, period_end,
                                     opening_balance, closing_balance, total_transactions, rent_transactions,
                                     bank_format, status)
        VALUES (?, ?, NULL, ?, ?, ?, ?, ?, ?, ?, ?, 'family_bank', 'active')
    """, (stmt_id, org_id, os.path.basename(pdf_path) + ' (reconciliation Oct 2026)', pdf_path,
          dates[0], dates[-1], float(result['opening_balance']), float(result['closing_balance']),
          len(txns), sum(1 for t in txns if t.txn_type == 'PAYBILL_CREDIT')))
    added = skipped = 0
    for t in txns:
        if t.reference and _norm(t.reference) in known:
            skipped += 1
            continue
        conn.execute("""
            INSERT INTO bank_transactions (id, statement_id, mpesa_ref, amount, txn_type, sender_name,
                                           txn_date, raw_text, unit_hint)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (generate_id('TXN'), stmt_id, t.reference, float(t.amount), t.txn_type, t.sender,
              _iso(t.transaction_date, t.raw_text), t.raw_text or '', t.unit_hint))
        added += 1
    return stmt_id, f'{os.path.basename(pdf_path)}: {dates[0]} → {dates[-1]}, {added} rows added, {skipped} already in system'


def add_verified_rows(conn, org_id, csv_path):
    """CSV columns: date (YYYY-MM-DD), reference, mpesa_code, payer, amount, unit, confidence, page, note."""
    with open(csv_path, newline='') as f:
        rows = [r for r in csv.DictReader(f) if (r.get('amount') or '').strip()]
    known = _known_refs(conn, org_id)
    dates = sorted(r['date'] for r in rows)
    stmt_id = generate_id('STMT')
    conn.execute("""
        INSERT INTO bank_statements (id, org_id, property_id, filename, file_path, period_start, period_end,
                                     total_transactions, rent_transactions, bank_format, status)
        VALUES (?, ?, NULL, ?, ?, ?, ?, ?, ?, 'national_bank', 'active')
    """, (stmt_id, org_id, os.path.basename(csv_path) + ' (Mowin lines checked by eye against the scanned copy)',
          csv_path, dates[0], dates[-1], len(rows), len(rows)))
    # The earlier National Bank scan (Mar–Apr) was read by the unreliable scanner, so its
    # reference codes can be wrong: also treat the same date + amount there as already present.
    same_day_amount = {}
    for r in conn.execute("""
        SELECT bt.id, bt.txn_date, bt.amount, bt.sender_name,
               EXISTS (SELECT 1 FROM payments p WHERE p.bank_txn_id = bt.id) AS assigned
        FROM bank_transactions bt JOIN bank_statements bs ON bs.id = bt.statement_id
        WHERE bs.org_id = ? AND bs.bank_format IN ('national_bank', 'scanned')""", (org_id,)):
        same_day_amount.setdefault((r['txn_date'], round(float(r['amount']), 2)), r)
    added = skipped = enriched = 0
    for r in rows:
        ref = (r.get('mpesa_code') or '').strip() or r['reference'].strip()
        amount_key = (r['date'], round(float(Decimal(r['amount'].replace(',', ''))), 2))
        raw = (f"National Bank {r['date']} ref {r['reference']} | {r['payer']} | KES {r['amount']} | "
               f"checked by eye against scanned statement page {r.get('page', '?')} | "
               f"confidence: {(r.get('confidence') or 'LOW').strip().upper()}"
               + (f" | {r['note']}" if r.get('note') else ''))
        if _norm(ref) in known or _norm(r['reference']) in known:
            skipped += 1
            continue
        twin = same_day_amount.get(amount_key)
        if twin:
            # Already in the system from the earlier scan, often with no payer name: add what the
            # checked reading says (name, likely unit) instead of a second copy of the same money.
            if not twin['assigned']:
                conn.execute("""UPDATE bank_transactions SET sender_name = COALESCE(NULLIF(sender_name, ''), ?),
                                unit_hint = COALESCE(unit_hint, ?), raw_text = COALESCE(raw_text, '') || ' || ' || ?
                                WHERE id = ?""", (r['payer'], (r.get('unit') or '').strip() or None, raw, twin['id']))
                enriched += 1
            else:
                skipped += 1
            continue
        conn.execute("""
            INSERT INTO bank_transactions (id, statement_id, mpesa_ref, amount, txn_type, sender_name,
                                           txn_date, raw_text, unit_hint)
            VALUES (?, ?, ?, ?, 'PAYBILL_CREDIT', ?, ?, ?, ?)
        """, (generate_id('TXN'), stmt_id, ref, float(Decimal(r['amount'].replace(',', ''))), r['payer'],
              r['date'], raw, (r.get('unit') or '').strip() or None))
        added += 1
    return stmt_id, (f'{os.path.basename(csv_path)}: {dates[0]} → {dates[-1]}, {added} rows added, '
                     f'{enriched} already in system and now named, {skipped} already assigned/present')


def add_missing_charges(conn, prop, periods):
    """Same rule as Generate charges: occupied units, rent + service, due per the lease (5th of that month)."""
    units = conn.execute("""
        SELECT u.* FROM units u WHERE u.property_id = ? AND u.status = 'occupied'
          AND EXISTS (SELECT 1 FROM tenants t WHERE t.unit_id = u.id AND t.status = 'active')
    """, (prop['id'],)).fetchall()
    from src.utils.due_dates import charge_due_date
    created = []
    for period in periods:
        due = charge_due_date(period, prop['rent_due_day'])
        for u in units:
            for ctype, amount in (('rent', u['monthly_rent']), ('service', u['service_charge'])):
                if not amount or float(amount) <= 0:
                    continue
                cur = conn.execute("""
                    INSERT OR IGNORE INTO rent_charges (id, property_id, unit_id, period, charge_type, amount, due_date)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                """, (generate_id('CHG'), prop['id'], u['id'], period, ctype, amount, due))
                if cur.rowcount:
                    created.append((u['unit_number'], period, ctype, float(amount)))
    conn.execute("INSERT INTO audit_log (action, entity_type, entity_id, details, user_id) VALUES (?, ?, ?, ?, ?)",
                 ('charges_generated', 'charge', ','.join(periods),
                  f"Reconciliation: {len(created)} missing rent/service charges for {periods[0]}–{periods[-1]}",
                  ACTOR))
    apply_unallocated_credit(conn, prop['id'])
    return created


def prepare(conn, prop, family, nbk, periods):
    log = []
    for p in family:
        log.append(add_family_statement(conn, prop['organization_id'], p)[1])
    for c in nbk:
        log.append(add_verified_rows(conn, prop['organization_id'], c)[1])
    charges = add_missing_charges(conn, prop, periods)
    log.append(f'{len(charges)} missing rent/service charges added '
               f'(KES {sum(c[3] for c in charges):,.0f}) for {periods[0]} → {periods[-1]}')
    return log, charges


def _balances(conn, pid):
    return {r['tenant_id']: dict(r) for r in conn.execute("""
        SELECT tb.tenant_id, tb.balance, t.name, u.unit_number FROM tenant_balances tb
        JOIN tenants t ON t.id = tb.tenant_id JOIN units u ON u.id = t.unit_id
        WHERE t.property_id = ? AND t.status = 'active'""", (pid,))}


def _possible_tenant(payer, tenants):
    """Best similar-name tenant for a payer the matcher could not place, or None.
    Only a hint for the reviewer — never pre-fills a decision."""
    import difflib, re
    words = [w for w in re.findall(r'[a-z]+', (payer or '').lower()) if len(w) > 2]
    best, best_score = None, 0.0
    for unit, name in tenants:
        tw = [w for w in re.findall(r'[a-z]+', name.lower()) if len(w) > 2]
        score = sum(max((difflib.SequenceMatcher(None, w, t).ratio() for t in tw), default=0) >= 0.8
                    for w in words)
        if score > best_score:
            best, best_score = (unit, name), score
    # Two similar words (one if the payer gave a single name): a shared first name alone is not a match.
    return best if words and best_score >= min(2, len(words)) else None


# ── Review sheet ─────────────────────────────────────────────────────────────

HEAD = ['Ref', 'Date', 'Amount (KES)', 'Payer on statement', 'Bank', 'Suggested unit', 'Suggested tenant',
        'Why suggested', 'DECISION', 'Your note']
SURE = ('unit_hint', 'name_match', 'sender_history', 'checked')


def write_review(conn, prop, out, log, charges, before):
    rows = get_unassigned_credits(conn, prop['id'], prop['organization_id'])
    stmt_bank = {r['id']: (r['bank_format'], r['filename']) for r in conn.execute(
        "SELECT id, bank_format, filename FROM bank_statements WHERE org_id = ?", (prop['organization_id'],))}
    unit_names = {r['unit_number'].upper(): r['name'] for r in conn.execute("""
        SELECT u.unit_number, t.name FROM units u JOIN tenants t ON t.unit_id = u.id AND t.status = 'active'
        WHERE u.property_id = ?""", (prop['id'],))}
    wb = Workbook()
    ws = wb.active
    ws.title = 'Payments to match'
    ws.append(HEAD)
    bold = Font(bold=True)
    for c in ws[1]:
        c.font = bold
    fill_sure = PatternFill('solid', fgColor='E8F5E9')
    fill_check = PatternFill('solid', fgColor='FFF8E1')
    for r in rows:
        bank, fname = stmt_bank.get(r['statement_id'], ('?', ''))
        verified = 'checked by eye' in (r.get('raw_text') or '')
        unit = r.get('suggested_unit_number') or ''
        source = r.get('suggestion_source') or ''
        if verified and r.get('unit_hint'):
            unit, source = r['unit_hint'], 'checked'
            raw = r.get('raw_text') or ''
            if 'confidence: HIGH' not in raw:
                source = 'checked_unsure'
        why = {'unit_hint': 'unit written in the payment', 'name_match': 'payer name matches tenant',
               'sender_history': 'this payer has paid for this unit before',
               'checked': 'National Bank line checked against the scan',
               'checked_unsure': 'National Bank line checked against the scan — payer not named, please confirm: '
                                 + ((r.get('raw_text') or '').rsplit('|', 1)[-1].strip())}.get(source, '')
        if not unit and not (bank in ('national_bank', 'scanned') and not verified):
            guess = _possible_tenant(r.get('sender_name'), list(unit_names.items()))
            if guess:
                unit, source = guess[0], 'possible'
                why = f'possible — similar name ({guess[1]}); please confirm'
        decision = unit if source in SURE else ''
        if not unit and bank in ('national_bank', 'scanned') and not verified:
            decision, why = 'IGNORE', 'National Bank line with no Mowin tenant name — probably law-firm income'
        ws.append([r.get('mpesa_ref') or '', r.get('txn_date') or '', float(r['amount']),
                   (r.get('sender_name') or '')[:60], 'National Bank' if bank in ('national_bank', 'scanned')
                   else 'Family Bank', unit, unit_names.get(unit.upper(), r.get('suggested_tenant_name') or ''),
                   why, decision, ''])
        for c in ws[ws.max_row]:
            c.fill = fill_sure if decision and decision != 'IGNORE' else fill_check
    for col, w in zip('ABCDEFGHIJ', (14, 11, 13, 34, 13, 13, 26, 40, 12, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'

    hs = wb.create_sheet('How to fill', 0)
    for line in [
        'How to check this sheet',
        '',
        'The "Payments to match" tab lists every payment the bank shows that is not yet linked to a tenant.',
        'Green rows: I am fairly sure — DECISION already holds the unit. Check it is right.',
        'Yellow rows: unsure, or no match — DECISION is blank or IGNORE.',
        '',
        'In DECISION write ONE of:',
        '   a unit number (e.g. 4C NBK, 1F)  → the payment is assigned to whoever lives there',
        '   IGNORE                           → not Mowin money (e.g. other law-firm income); hidden from the list',
        '   leave blank                      → not sure yet; it stays unassigned (e.g. a new tenant with no profile yet)',
        '',
        'Use "Your note" for anything I should know. Send the sheet back and I apply it, with a backup first.',
        'No SMS goes to tenants for these old payments.',
        '',
        'What was added before this sheet was made:',
    ] + ['   ' + l for l in log]:
        hs.append([line])
    hs['A1'].font = Font(bold=True, size=13)
    hs.column_dimensions['A'].width = 110

    cs = wb.create_sheet('Charges added')
    cs.append(['Unit', 'Month', 'Type', 'Amount (KES)'])
    for c in cs[1]:
        c.font = bold
    for c in charges:
        cs.append(list(c))

    bs = wb.create_sheet('Balances')
    bs.append(['Unit', 'Tenant (old record)', 'Balance before (KES)', 'After charges added (KES)',
               'Green matches if approved (KES)', 'Change (KES)'])
    for c in bs[1]:
        c.font = bold
    after = _balances(conn, prop['id'])
    green = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[8] and row[8] != 'IGNORE':
            green[row[8].upper()] = green.get(row[8].upper(), 0) + row[2]
    for tid, a in sorted(after.items(), key=lambda kv: kv[1]['unit_number']):
        b = before.get(tid, {}).get('balance') or 0
        projected = a['balance'] - green.get(a['unit_number'].upper(), 0)
        bs.append([a['unit_number'], a['name'], b, a['balance'], projected, projected - b])
    for sheet, cols in ((bs, 'BCDEF'), (cs, 'ABCD')):
        sheet.column_dimensions['B'].width = 28
    for col in 'CDEF':
        bs.column_dimensions[col].width = 22
    wb.save(out)
    return len(rows)


# ── Apply ────────────────────────────────────────────────────────────────────

def apply_decisions(conn, prop, review_path):
    ws = load_workbook(review_path)['Payments to match']
    pending = {}
    for r in get_unassigned_credits(conn, prop['id'], prop['organization_id']):
        pending.setdefault((_norm(r['mpesa_ref']), round(float(r['amount']), 2)), r)
    units = {r['unit_number'].upper(): r for r in conn.execute(
        "SELECT * FROM units WHERE property_id = ?", (prop['id'],))}
    assigned = ignored = 0
    problems = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        ref, day, amount, payer, _bank, _su, _st, why, decision, note = (list(row) + [None] * 10)[:10]
        decision = (str(decision or '')).strip()
        if not decision:
            continue
        txn = pending.get((_norm(ref), round(float(amount), 2)))
        if not txn:
            problems.append(f'{ref} KES {amount}: not found unassigned (already assigned, or ref changed) — skipped')
            continue
        if decision.upper() == 'IGNORE':
            conn.execute("""UPDATE bank_transactions SET ignored = 1, ignored_at = CURRENT_TIMESTAMP,
                            ignored_reason = ? WHERE id = ?""",
                         ((note or why or 'Not Mowin (reconciliation Oct 2026)')[:200], txn['id']))
            ignored += 1
            continue
        unit = units.get(decision.upper())
        if not unit:
            problems.append(f'{ref}: unit "{decision}" does not exist — skipped')
            continue
        reason = f"Reconciliation Oct 2026 — {why or 'matched by review'}" + (f" — {note}" if note else '')
        payment_id = generate_id('PAY')
        conn.execute("""
            INSERT INTO payments (id, property_id, unit_id, bank_txn_id, statement_id, amount, payment_date,
                                  assignment_type, assignment_reason, assigned_by)
            VALUES (?, ?, ?, ?, ?, ?, ?, 'manual', ?, ?)
        """, (payment_id, prop['id'], unit['id'], txn['id'], txn['statement_id'], txn['amount'],
              txn['txn_date'] or '', reason, ACTOR))
        allocate_payment(conn, payment_id, unit['id'], txn['amount'])
        conn.execute("INSERT INTO audit_log (action, entity_type, entity_id, details, user_id) VALUES (?, ?, ?, ?, ?)",
                     ('payment_manual_assigned', 'payment', payment_id,
                      f"Ref: {ref} | Amount: KES {float(amount):,.2f} | Unit: {unit['unit_number']} | Reason: {reason}",
                      ACTOR))
        assigned += 1
        pending.pop((_norm(ref), round(float(amount), 2)), None)
    return assigned, ignored, problems


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('mode', choices=['review', 'apply'])
    ap.add_argument('--property', default='PROP-666759B2')
    ap.add_argument('--family', nargs='*', default=[])
    ap.add_argument('--nbk', nargs='*', default=[])
    ap.add_argument('--periods', nargs='*', default=DEFAULT_PERIODS)
    ap.add_argument('--out', default='reconciliation_review.xlsx')
    ap.add_argument('--review')
    ap.add_argument('--execute', action='store_true')
    a = ap.parse_args()
    print('Database:', os.environ.get('DATABASE_PATH'))
    try:
        with get_connection() as conn:
            prop = conn.execute("SELECT * FROM properties WHERE id = ?", (a.property,)).fetchone()
            if not prop:
                raise SystemExit('No such property')
            if conn.execute("SELECT 1 FROM audit_log WHERE user_id = ? AND action = 'charges_generated'",
                            (ACTOR,)).fetchone():
                raise SystemExit('This reconciliation has already been applied to this database.')
            before = _balances(conn, prop['id'])
            log, charges = prepare(conn, prop, a.family, a.nbk, a.periods)
            for line in log:
                print(' ', line)
            if a.mode == 'review':
                n = write_review(conn, prop, a.out, log, charges, before)
                print(f'Review sheet: {a.out} ({n} payments to match)')
            else:
                if not a.review:
                    raise SystemExit('--review reviewed.xlsx is required')
                assigned, ignored, problems = apply_decisions(conn, prop, a.review)
                print(f'Assigned {assigned}, ignored {ignored}.')
                for p in problems:
                    print('  !', p)
                after = _balances(conn, prop['id'])
                total_b = sum((v['balance'] or 0) for v in before.values())
                total_a = sum((v['balance'] or 0) for v in after.values())
                print(f'Total owed by current tenants: KES {total_b:,.0f} → KES {total_a:,.0f}')
                if not a.execute:
                    raise DryRun()
    except DryRun:
        print('Dry run — nothing saved. Add --execute to apply.')


if __name__ == '__main__':
    main()
